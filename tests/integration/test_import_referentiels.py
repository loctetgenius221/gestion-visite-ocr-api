"""Import en masse des référentiels : modèles, prévisualisation, application.

Le fil conducteur est le même partout : **le plan fait foi**. Ce que la
prévisualisation annonce est ce que l'application exécute, et rien ne s'écrit
tant qu'une seule ligne cloche. Les tests vérifient donc autant ce qui est créé
que ce qui ne l'est **pas** — un import qui écrit à moitié serait pire qu'un
import qui échoue, parce qu'on ne saurait plus où il s'est arrêté.
"""

from __future__ import annotations

import io

import pytest
from httpx import AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.referentiel import Agent, Service

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def classeur(entetes: list[str], lignes: list[list[object]]) -> bytes:
    """Fabrique un `.xlsx` en mémoire, comme le ferait un administrateur."""
    wb = Workbook()
    feuille = wb.active
    feuille.append(entetes)
    for ligne in lignes:
        feuille.append(ligne)
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


def services_xlsx(lignes: list[list[object]]) -> bytes:
    return classeur(["Code", "Nom", "Étage", "Code parent"], lignes)


def agents_xlsx(lignes: list[list[object]]) -> bytes:
    return classeur(["Nom", "Code service", "Fonction", "Bureau"], lignes)


def envoi(contenu: bytes, nom: str = "import.xlsx") -> dict:
    return {"files": {"fichier": (nom, contenu, XLSX)}}


class TestModeles:
    async def test_le_modele_services_est_un_classeur_exploitable(
        self, client: AsyncClient, org_admin_headers: dict[str, str]
    ) -> None:
        reponse = await client.get("/services/import/template", headers=org_admin_headers)

        assert reponse.status_code == 200, reponse.text
        assert reponse.headers["content-type"].startswith(XLSX)
        assert "modele-services.xlsx" in reponse.headers["content-disposition"]

        wb = load_workbook(io.BytesIO(reponse.content))
        assert wb.sheetnames == ["Services", "Notice"]
        assert [c.value for c in wb["Services"][1]] == ["Code", "Nom", "Étage", "Code parent"]

    async def test_la_feuille_de_saisie_ne_contient_aucun_exemple(
        self, client: AsyncClient, org_admin_headers: dict[str, str]
    ) -> None:
        """Un exemple laissé dans la grille finirait importé comme une vraie ligne."""
        reponse = await client.get("/agents/import/template", headers=org_admin_headers)

        wb = load_workbook(io.BytesIO(reponse.content))
        assert wb["Agents"].max_row == 1  # la ligne d'en-tête, et rien d'autre
        # Les exemples existent, mais sur la notice.
        notice = "\n".join(
            str(c.value) for rangee in wb["Notice"].iter_rows() for c in rangee if c.value
        )
        assert "DRH" in notice

    async def test_les_colonnes_sont_au_format_texte(
        self, client: AsyncClient, org_admin_headers: dict[str, str]
    ) -> None:
        """Sans ça, le code de service « 007 » devient le nombre 7 à la saisie.

        Le format est porté par la colonne, pas par des cellules pré-créées : il
        s'applique donc à tout ce que l'administrateur saisira, sans laisser mille
        lignes vides dans le fichier.
        """
        reponse = await client.get("/services/import/template", headers=org_admin_headers)

        wb = load_workbook(io.BytesIO(reponse.content))
        colonnes = wb["Services"].column_dimensions
        assert [colonnes[lettre].number_format for lettre in "ABCD"] == ["@"] * 4

    async def test_un_agent_de_controle_na_pas_acces_aux_modeles(
        self, client: AsyncClient, auth_headers: dict[str, str]
    ) -> None:
        reponse = await client.get("/services/import/template", headers=auth_headers)

        assert reponse.status_code == 403


class TestPrevisualisationServices:
    async def test_le_plan_annonce_les_creations_sans_rien_ecrire(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        contenu = services_xlsx([["FIN", "Direction Financière", "1er", ""]])

        reponse = await client.post(
            "/services/import/preview", headers=org_admin_headers, **envoi(contenu)
        )

        assert reponse.status_code == 200, reponse.text
        plan = reponse.json()
        assert plan["applique"] is False
        assert (plan["total"], plan["a_creer"], plan["erreurs"]) == (1, 1, 0)
        assert plan["items"][0]["statut"] == "create"
        assert plan["items"][0]["ligne"] == 2

        apres = await client.get("/services", headers=org_admin_headers)
        assert "FIN" not in [s["code"] for s in apres.json()]

    async def test_le_plan_nomme_lorganisation_visee(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        """Un import massif dans la mauvaise structure ne doit pas se découvrir après."""
        reponse = await client.post(
            "/services/import/preview",
            headers=org_admin_headers,
            **envoi(services_xlsx([["FIN", "Direction Financière", "", ""]])),
        )

        assert reponse.json()["organization_nom"] == seeded["organization"].nom
        assert reponse.json()["organization_id"] == str(seeded["organization"].id)

    async def test_un_service_deja_present_ressort_inchange(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        contenu = services_xlsx([["DRH", "Un autre libellé", "", ""]])

        reponse = await client.post(
            "/services/import/preview", headers=org_admin_headers, **envoi(contenu)
        )

        item = reponse.json()["items"][0]
        assert item["statut"] == "unchanged"
        assert item["entity_id"] == str(seeded["service"].id)

    async def test_avec_mise_a_jour_le_meme_service_ressort_modifie(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        contenu = services_xlsx([["DRH", "Un autre libellé", "", ""]])

        reponse = await client.post(
            "/services/import/preview?mettre_a_jour=true",
            headers=org_admin_headers,
            **envoi(contenu),
        )

        item = reponse.json()["items"][0]
        assert item["statut"] == "update"
        assert item["changements"]["name"]["après"] == "Un autre libellé"

    async def test_une_cellule_vide_nefface_pas_la_valeur_en_base(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        """« Je ne renseigne pas » n'est pas « efface ce qui est en base »."""
        contenu = services_xlsx([["DRH", "Direction des RH", "", ""]])

        reponse = await client.post(
            "/services/import/preview?mettre_a_jour=true",
            headers=org_admin_headers,
            **envoi(contenu),
        )

        # L'étage du jeu de test vaut « 2e étage » : la cellule vide le laisse.
        assert reponse.json()["items"][0]["statut"] == "unchanged"


class TestHierarchieDesServices:
    async def test_un_parent_defini_plus_bas_dans_le_fichier_est_accepte(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        """L'ordre des lignes est l'affaire du serveur, pas de l'administrateur."""
        contenu = services_xlsx(
            [
                ["ENFANT", "Bureau des enfants perdus", "", "PARENT"],
                ["PARENT", "Direction parente", "", ""],
            ]
        )

        reponse = await client.post(
            "/services/import/preview", headers=org_admin_headers, **envoi(contenu)
        )

        assert reponse.json()["erreurs"] == 0
        assert reponse.json()["a_creer"] == 2

    async def test_un_parent_introuvable_est_signale_sur_sa_ligne(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        contenu = services_xlsx([["FIN", "Direction Financière", "", "FANTOME"]])

        reponse = await client.post(
            "/services/import/preview", headers=org_admin_headers, **envoi(contenu)
        )

        item = reponse.json()["items"][0]
        assert item["statut"] == "error"
        assert item["error_code"] == "IMPORT_PARENT_UNKNOWN"
        assert item["ligne"] == 2

    async def test_un_service_ne_peut_pas_etre_son_propre_parent(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        contenu = services_xlsx([["BOUCLE", "Service bouclé", "", "BOUCLE"]])

        reponse = await client.post(
            "/services/import/preview", headers=org_admin_headers, **envoi(contenu)
        )

        assert reponse.json()["items"][0]["error_code"] == "SERVICE_HIERARCHY_CYCLE"

    async def test_un_cycle_entre_deux_lignes_du_fichier_est_refuse(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        """`_build_tree` reparcourt la hiérarchie à chaque listing : une boucle
        y tournerait sans fin."""
        contenu = services_xlsx([["A", "Service A", "", "B"], ["B", "Service B", "", "A"]])

        reponse = await client.post(
            "/services/import/preview", headers=org_admin_headers, **envoi(contenu)
        )

        codes = {
            item["error_code"] for item in reponse.json()["items"] if item["statut"] == "error"
        }
        assert codes == {"SERVICE_HIERARCHY_CYCLE"}
        assert reponse.json()["erreurs"] >= 1

    async def test_un_cycle_qui_se_referme_par_la_base_est_refuse(
        self,
        client: AsyncClient,
        org_admin_headers: dict[str, str],
        session: AsyncSession,
        seeded: dict,
    ) -> None:
        """Le graphe ne peut pas être celui du seul fichier : la moitié du cycle est en base."""
        parent = seeded["service"]
        enfant = seeded["other_service"]
        enfant.parent_id = parent.id
        await session.commit()

        # Le fichier veut rattacher le parent à son propre enfant.
        contenu = services_xlsx([[parent.code, parent.name, "", enfant.code]])

        reponse = await client.post(
            "/services/import/preview?mettre_a_jour=true",
            headers=org_admin_headers,
            **envoi(contenu),
        )

        assert reponse.json()["items"][0]["error_code"] == "SERVICE_HIERARCHY_CYCLE"

    async def test_lapplication_cree_larbre_dans_le_bon_ordre(
        self,
        client: AsyncClient,
        org_admin_headers: dict[str, str],
        session: AsyncSession,
        seeded: dict,
    ) -> None:
        lignes = [
            {"code": "PETIT", "nom": "Sous-bureau", "code_parent": "MOYEN"},
            {"code": "MOYEN", "nom": "Bureau", "code_parent": "GRAND"},
            {"code": "GRAND", "nom": "Direction"},
        ]

        reponse = await client.post(
            "/services/import", headers=org_admin_headers, json={"lignes": lignes}
        )

        assert reponse.status_code == 200, reponse.text
        assert reponse.json()["applique"] is True

        crees = {
            s.code: s
            for s in (await session.execute(select(Service))).scalars().all()
            if s.code in {"PETIT", "MOYEN", "GRAND"}
        }
        assert crees["GRAND"].parent_id is None
        assert crees["MOYEN"].parent_id == crees["GRAND"].id
        assert crees["PETIT"].parent_id == crees["MOYEN"].id


class TestApplicationServices:
    async def test_une_seule_ligne_en_erreur_annule_tout_limport(
        self,
        client: AsyncClient,
        org_admin_headers: dict[str, str],
        session: AsyncSession,
        seeded: dict,
    ) -> None:
        """Le tout-ou-rien est le contrat : un import à moitié appliqué serait pire."""
        lignes = [
            {"code": "BON", "nom": "Service valide"},
            {"code": "MAUVAIS", "nom": "Service bancal", "code_parent": "FANTOME"},
        ]

        reponse = await client.post(
            "/services/import", headers=org_admin_headers, json={"lignes": lignes}
        )

        assert reponse.status_code == 422
        assert reponse.json()["error_code"] == "IMPORT_PLAN_REJECTED"
        codes = [s.code for s in (await session.execute(select(Service))).scalars().all()]
        assert "BON" not in codes

    async def test_les_services_importes_apparaissent_dans_le_listing(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        reponse = await client.post(
            "/services/import",
            headers=org_admin_headers,
            json={"lignes": [{"code": "FIN", "nom": "Direction Financière", "etage": "1er"}]},
        )

        assert reponse.status_code == 200, reponse.text
        listing = await client.get("/services", headers=org_admin_headers)
        assert "FIN" in [s["code"] for s in listing.json()]

    async def test_rejouer_le_meme_import_ne_cree_pas_de_doublon(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        corps = {"lignes": [{"code": "FIN", "nom": "Direction Financière"}]}

        premier = await client.post("/services/import", headers=org_admin_headers, json=corps)
        second = await client.post("/services/import", headers=org_admin_headers, json=corps)

        assert premier.json()["a_creer"] == 1
        assert second.json()["a_creer"] == 0
        assert second.json()["inchangees"] == 1

    async def test_un_code_en_double_dans_le_fichier_est_refuse(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        contenu = services_xlsx([["FIN", "Premier libellé", "", ""], ["fin", "Second", "", ""]])

        reponse = await client.post(
            "/services/import/preview", headers=org_admin_headers, **envoi(contenu)
        )

        items = reponse.json()["items"]
        assert items[0]["statut"] == "create"
        assert items[1]["error_code"] == "IMPORT_DUPLICATE_IN_FILE"

    async def test_limport_est_journalise_deux_fois(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        """Une trace par entité pour « qui a créé ça ? », une de synthèse pour « ce jour-là ? »."""
        await client.post(
            "/services/import",
            headers=org_admin_headers,
            json={"lignes": [{"code": "FIN", "nom": "Direction Financière"}]},
        )

        journal = await client.get(
            "/audit-logs?page_size=100", headers=org_admin_headers
        )
        actions = [entree["action"] for entree in journal.json()["items"]]
        assert "referentiel.imported" in actions
        assert "service.created" in actions


class TestImportDesAgents:
    async def test_un_agent_est_rattache_par_le_code_de_son_service(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        reponse = await client.post(
            "/agents/import",
            headers=org_admin_headers,
            json={
                "lignes": [
                    {
                        "nom": "Fatou SOW",
                        "code_service": "DRH",
                        "fonction": "Chef de bureau",
                        "bureau": "B105",
                    }
                ]
            },
        )

        assert reponse.status_code == 200, reponse.text
        listing = await client.get("/agents", headers=org_admin_headers)
        fatou = next(a for a in listing.json() if a["name"] == "Fatou SOW")
        assert fatou["role"] == "Chef de bureau"
        assert fatou["organization_id"] == str(seeded["organization"].id)

    async def test_un_code_de_service_inconnu_est_signale_sur_sa_ligne(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        """L'import d'agents ne crée jamais de service : une faute de frappe en
        fabriquerait une direction fantôme."""
        contenu = agents_xlsx([["Fatou SOW", "INEXISTANT", "", ""]])

        reponse = await client.post(
            "/agents/import/preview", headers=org_admin_headers, **envoi(contenu)
        )

        item = reponse.json()["items"][0]
        assert item["statut"] == "error"
        assert item["error_code"] == "SERVICE_NOT_FOUND"

    async def test_un_service_archive_est_refuse(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        """Même règle qu'à la création unitaire : la personne serait sélectionnable
        sans service visible côté mobile."""
        # Un service portant des agents actifs ne peut pas être archivé : on en
        # crée un vide pour isoler la règle testée.
        vide = await client.post(
            "/services", headers=org_admin_headers, json={"name": "Service clos", "code": "CLOS"}
        )
        archivage = await client.patch(
            f"/services/{vide.json()['id']}/status",
            headers=org_admin_headers,
            json={"status": "archived"},
        )
        assert archivage.status_code == 200, archivage.text

        reponse = await client.post(
            "/agents/import/preview",
            headers=org_admin_headers,
            **envoi(agents_xlsx([["Nouveau Venu", "CLOS", "", ""]])),
        )

        assert reponse.json()["items"][0]["error_code"] == "ARCHIVED_REFERENTIEL"

    async def test_deux_fois_la_meme_personne_dans_le_meme_service(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        contenu = agents_xlsx(
            [["Fatou SOW", "DRH", "", ""], ["fatou sow", "DRH", "Autre poste", ""]]
        )

        reponse = await client.post(
            "/agents/import/preview", headers=org_admin_headers, **envoi(contenu)
        )

        assert reponse.json()["items"][1]["error_code"] == "IMPORT_DUPLICATE_IN_FILE"

    async def test_une_personne_deja_presente_est_mise_a_jour_sur_demande(
        self,
        client: AsyncClient,
        org_admin_headers: dict[str, str],
        session: AsyncSession,
        seeded: dict,
    ) -> None:
        corps = {
            "mettre_a_jour": True,
            "lignes": [
                {"nom": "Aminata Diallo", "code_service": "DRH", "fonction": "DRH adjointe"}
            ],
        }

        reponse = await client.post("/agents/import", headers=org_admin_headers, json=corps)

        assert reponse.json()["a_mettre_a_jour"] == 1
        agent = (
            await session.execute(select(Agent).where(Agent.id == seeded["agent"].id))
        ).scalar_one()
        await session.refresh(agent)
        assert agent.role == "DRH adjointe"


class TestPerimetreEtOrganisation:
    async def test_ladmin_de_plateforme_doit_nommer_lorganisation(
        self, client: AsyncClient, admin_headers: dict[str, str], seeded: dict
    ) -> None:
        """Un service appartient toujours à une structure : sans elle, il n'apparaîtrait
        dans la liste d'aucun poste de contrôle."""
        reponse = await client.post(
            "/services/import",
            headers=admin_headers,
            json={"lignes": [{"code": "FIN", "nom": "Direction Financière"}]},
        )

        assert reponse.status_code == 400
        assert reponse.json()["error_code"] == "ORGANIZATION_REQUIRED"

    async def test_une_organisation_inconnue_repond_404_et_non_500(
        self, client: AsyncClient, admin_headers: dict[str, str], seeded: dict
    ) -> None:
        """Sans ce contrôle, l'identifiant partait jusqu'à la contrainte de clé étrangère."""
        reponse = await client.post(
            "/services/import",
            headers=admin_headers,
            json={
                "organization_id": "00000000-0000-0000-0000-000000000000",
                "lignes": [{"code": "FIN", "nom": "Direction Financière"}],
            },
        )

        assert reponse.status_code == 404
        assert reponse.json()["error_code"] == "ORGANIZATION_NOT_FOUND"

    async def test_ladmin_de_plateforme_importe_dans_la_structure_designee(
        self, client: AsyncClient, admin_headers: dict[str, str], voisine: dict
    ) -> None:
        reponse = await client.post(
            "/services/import",
            headers=admin_headers,
            json={
                "organization_id": str(voisine["organization"].id),
                "lignes": [{"code": "QUAI", "nom": "Direction des Quais"}],
            },
        )

        assert reponse.status_code == 200, reponse.text
        assert reponse.json()["organization_nom"] == voisine["organization"].nom

        chez_voisine = await client.get(
            f"/services?organization_id={voisine['organization'].id}", headers=admin_headers
        )
        assert "QUAI" in [s["code"] for s in chez_voisine.json()]

    async def test_un_org_admin_ne_peut_pas_importer_ailleurs(
        self, client: AsyncClient, org_admin_headers: dict[str, str], voisine: dict
    ) -> None:
        reponse = await client.post(
            "/services/import",
            headers=org_admin_headers,
            json={
                "organization_id": str(voisine["organization"].id),
                "lignes": [{"code": "QUAI", "nom": "Direction des Quais"}],
            },
        )

        assert reponse.status_code == 403
        assert reponse.json()["error_code"] == "CROSS_ORGANIZATION_FORBIDDEN"

    async def test_un_agent_de_controle_ne_peut_pas_importer(
        self, client: AsyncClient, auth_headers: dict[str, str], seeded: dict
    ) -> None:
        reponse = await client.post(
            "/services/import",
            headers=auth_headers,
            json={"lignes": [{"code": "FIN", "nom": "Direction Financière"}]},
        )

        assert reponse.status_code == 403


class TestFichiersMalformes:
    async def test_un_vrai_csv_est_refuse(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        reponse = await client.post(
            "/services/import/preview",
            headers=org_admin_headers,
            **envoi(b"Code;Nom\nFIN;Finances", "annuaire.csv"),
        )

        assert reponse.status_code == 400
        assert reponse.json()["error_code"] == "UNSUPPORTED_IMPORT_FILE"
        # Le message nomme ce qui a été reçu : « réenregistrez au bon format » sans
        # dire lequel on a lu n'aide personne.
        assert "CSV" in reponse.json()["message"]

    async def test_lextension_ne_suffit_pas_le_contenu_est_verifie(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        """Même règle que pour le logo d'organisation : on lit les octets, pas le nom."""
        reponse = await client.post(
            "/services/import/preview",
            headers=org_admin_headers,
            **envoi(b"Code;Nom\nFIN;Finances", "annuaire.xlsx"),
        )

        assert reponse.status_code == 400
        assert reponse.json()["error_code"] == "UNSUPPORTED_IMPORT_FILE"

    async def test_un_classeur_valide_est_accepte_meme_mal_nomme(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        """La règle du contenu joue dans les deux sens.

        Un client HTTP qui ignore `Content-Disposition` enregistre le modèle sous
        un nom de son cru — vu en `export.csv`. Refuser là-dessus ferait rejeter
        notre propre classeur, intact, pour une raison qui ne regarde pas son
        contenu.
        """
        reponse = await client.post(
            "/services/import/preview",
            headers=org_admin_headers,
            **envoi(services_xlsx([["FIN", "Direction Financière", "", ""]]), "export.csv"),
        )

        assert reponse.status_code == 200, reponse.text
        assert reponse.json()["a_creer"] == 1

    async def test_le_modele_telecharge_est_reimportable_tel_quel(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        """L'aller-retour complet : ce que la route sert, la route le relit.

        Le modèle est vide, donc l'import ne porte sur rien — c'est le passage des
        en-têtes et du format qui est vérifié ici, bout à bout.
        """
        modele = await client.get("/services/import/template", headers=org_admin_headers)
        assert modele.status_code == 200

        reponse = await client.post(
            "/services/import/preview",
            headers=org_admin_headers,
            **envoi(modele.content, "modele-services.xlsx"),
        )

        assert reponse.status_code == 200, reponse.text
        assert reponse.json()["total"] == 0

    async def test_des_colonnes_manquantes_disent_ce_qui_a_ete_lu(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        """Sans la liste des colonnes lues, on ne devine pas qu'une ligne de titre
        s'est glissée au-dessus des en-têtes."""
        contenu = classeur(["Intitulé", "Bâtiment"], [["Finances", "A"]])

        reponse = await client.post(
            "/services/import/preview", headers=org_admin_headers, **envoi(contenu)
        )

        assert reponse.status_code == 400
        details = reponse.json()["details"]
        assert details["colonnes_manquantes"] == ["Code"]
        assert "Bâtiment" in details["colonnes_lues"]

    async def test_les_entetes_tolerent_casse_accents_et_separateurs(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        contenu = classeur(
            ["CODE", "libelle", "niveau", "code_parent"], [["FIN", "Finances", "1er", ""]]
        )

        reponse = await client.post(
            "/services/import/preview", headers=org_admin_headers, **envoi(contenu)
        )

        assert reponse.status_code == 200, reponse.text
        assert reponse.json()["a_creer"] == 1

    async def test_une_formule_non_calculee_est_signalee(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        """Lue telle quelle, la cellule paraîtrait vide et la ligne partirait amputée."""
        contenu = services_xlsx([["=CONCATENER(A1;B1)", "Direction Financière", "", ""]])

        reponse = await client.post(
            "/services/import/preview", headers=org_admin_headers, **envoi(contenu)
        )

        assert reponse.status_code == 400
        assert reponse.json()["error_code"] == "IMPORT_UNCOMPUTED_FORMULA"

    async def test_un_fichier_trop_long_est_refuse(
        self, client: AsyncClient, admin_headers: dict[str, str], seeded: dict
    ) -> None:
        """Symétrique du plafond d'export : l'import écrit, là où l'export ne fait que lire."""
        await client.put(
            "/settings", headers=admin_headers, json={"referentiel_import_max_rows": 10}
        )
        contenu = services_xlsx([[f"S{i}", f"Service {i}", "", ""] for i in range(15)])

        reponse = await client.post(
            "/services/import/preview", headers=admin_headers, **envoi(contenu)
        )

        assert reponse.status_code == 422
        assert reponse.json()["error_code"] == "IMPORT_TOO_MANY_ROWS"

    async def test_une_valeur_trop_longue_est_signalee_sur_sa_ligne(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        """Rejetée sur sa ligne plutôt que tronquée en silence : la cellule est corrigeable."""
        contenu = services_xlsx([["FIN", "N" * 250, "", ""]])

        reponse = await client.post(
            "/services/import/preview", headers=org_admin_headers, **envoi(contenu)
        )

        item = reponse.json()["items"][0]
        assert item["error_code"] == "IMPORT_VALUE_TOO_LONG"

    async def test_un_parametre_de_requete_inconnu_reste_refuse(
        self, client: AsyncClient, org_admin_headers: dict[str, str], seeded: dict
    ) -> None:
        """Les nouvelles routes héritent du garde-fou posé sur les routeurs."""
        reponse = await client.post(
            "/services/import/preview?maj=true",
            headers=org_admin_headers,
            **envoi(services_xlsx([["FIN", "Finances", "", ""]])),
        )

        assert reponse.status_code == 400
        assert reponse.json()["error_code"] == "VALIDATION_ERROR"


class TestNeutralisationDesFormules:
    @pytest.mark.parametrize("prefixe", ["=", "+", "-", "@"])
    async def test_un_nom_importe_ne_devient_pas_une_formule_a_lexport(
        self,
        client: AsyncClient,
        org_admin_headers: dict[str, str],
        admin_headers: dict[str, str],
        session: AsyncSession,
        seeded: dict,
        prefixe: str,
    ) -> None:
        """Le nom fait l'aller-retour : importé ici, il ressort dans l'export du
        registre — le fichier que l'administrateur ouvre dans son tableur."""
        nom = f"{prefixe}HYPERLINK(\"http://x\")"
        await client.post(
            "/agents/import",
            headers=org_admin_headers,
            json={"lignes": [{"nom": nom, "code_service": "DRH"}]},
        )

        # La base garde le nom tel quel : c'est peut-être un vrai nom.
        agents = (await session.execute(select(Agent))).scalars().all()
        assert any(a.name == nom for a in agents)

        export = await client.get("/visits/export?format=csv", headers=admin_headers)
        contenu = export.content.decode("utf-8-sig")
        assert f";{prefixe}HYPERLINK" not in contenu
