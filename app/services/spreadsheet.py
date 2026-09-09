"""Lecture et écriture de classeurs `.xlsx`, pour les imports de référentiels.

Le format est imposé plutôt que subi. Un CSV laisse à l'utilisateur trois façons
de casser son fichier sans s'en apercevoir — il l'ouvre dans Excel, enregistre, et
le voilà en ANSI, séparé par des virgules, avec le code de service « 007 » devenu
le nombre 7. Un `.xlsx` verrouille l'encodage, le séparateur et le type des
cellules : ce que l'administrateur a saisi est ce que le serveur relira.

Rien ici ne connaît les services ni les agents : ce module ramène des lignes de
`dict[str, str]` et sait fabriquer un classeur modèle. Le sens des colonnes est
l'affaire de `referentiel_import_service`.
"""

from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from app.core.config import settings
from app.core.errors import (
    FileTooLargeError,
    ImportColumnsMissingError,
    ImportTooManyRowsError,
    UncomputedFormulaError,
    UnsupportedImportFileError,
)

# Apostrophes typographiques : un même libellé mélange U+2019 et U+0027 selon
# qu'il a été tapé dans Word, dans Excel ou collé depuis un PDF.
APOSTROPHES = str.maketrans({"’": "'", "‘": "'", "ʼ": "'", "`": "'"})

# Caractères qui font d'une cellule une **formule** à l'ouverture du tableur.
# Une valeur qui commence par l'un d'eux est préfixée d'une apostrophe avant
# d'être écrite : voir `neutraliser_formule`.
AMORCES_FORMULE = ("=", "+", "-", "@", "\t", "\r")


def sans_accents(texte: str) -> str:
    decompose = unicodedata.normalize("NFKD", texte.translate(APOSTROPHES))
    return "".join(c for c in decompose if not unicodedata.combining(c))


def normalise(texte: str) -> str:
    """Clé de comparaison : sans accent, sans casse, espaces normalisés."""
    return " ".join(sans_accents(texte).casefold().split())


def normalise_entete(texte: str) -> str:
    """Même chose pour un en-tête de colonne, séparateurs ramenés à l'espace.

    « Code_Parent », « code parent » et « CODE-PARENT » désignent la même colonne :
    l'administrateur qui remplit le modèle n'a pas à en respecter la typographie.
    """
    return " ".join(re.sub(r"[_\-./]+", " ", sans_accents(texte)).casefold().split())


def neutraliser_formule(valeur: str) -> str:
    """Préfixe d'une apostrophe une valeur que le tableur prendrait pour une formule.

    Un nom d'agent commençant par `=`, `+`, `-` ou `@` devient une formule à
    l'ouverture du fichier. Le risque n'est pas théorique ici : les noms importés
    **ressortent** par l'export du registre (`visits_to_csv`), dans un fichier
    qu'un administrateur ouvre dans son tableur. La donnée est donc laissée
    intacte en base — c'est peut-être un vrai nom — et neutralisée à l'écriture,
    seul endroit où elle devient dangereuse.
    """
    return f"'{valeur}" if valeur.startswith(AMORCES_FORMULE) else valeur


def cellule_en_texte(valeur: object) -> str:
    """Rend une cellule sous forme de chaîne exploitable.

    Excel typant tout ce qu'il peut, un code de service saisi « 01 » revient en
    entier `1`, et une date mal devinée en `datetime`. Un `float` entier est rendu
    sans sa décimale — sans quoi le code `2024` deviendrait « 2024.0 » et ne
    correspondrait plus à rien en base.
    """
    if valeur is None:
        return ""
    if isinstance(valeur, bool):
        return "oui" if valeur else "non"
    if isinstance(valeur, float) and valeur.is_integer():
        return str(int(valeur))
    if isinstance(valeur, datetime | date):
        return valeur.isoformat()
    # L'apostrophe de neutralisation posée à l'écriture est retirée à la relecture,
    # pour qu'un aller-retour modèle → import rende la valeur d'origine.
    texte = str(valeur).strip()
    if texte.startswith("'") and texte[1:2] in AMORCES_FORMULE:
        texte = texte[1:]
    return texte


@dataclass(frozen=True, slots=True)
class Colonne:
    """Une colonne du modèle : son identité interne, son libellé, ses alias."""

    cle: str
    libelle: str
    obligatoire: bool
    description: str
    largeur: int = 28
    alias: tuple[str, ...] = ()
    exemple: str = ""

    @property
    def alias_normalises(self) -> tuple[str, ...]:
        """Libellés acceptés en en-tête, le libellé officiel compris."""
        return tuple({normalise_entete(x) for x in (self.libelle, self.cle, *self.alias)})


def resoudre_colonnes(entetes: list[str | None], colonnes: tuple[Colonne, ...]) -> dict[str, str]:
    """Associe chaque colonne attendue à l'en-tête réellement présent.

    Lève `ImportColumnsMissingError` si une colonne obligatoire manque, en listant
    ce qui a été lu : sans cette liste, l'administrateur ne peut pas deviner que
    son fichier a une ligne de titre au-dessus des en-têtes.
    """
    presentes = {normalise_entete(e): e for e in entetes if e}
    resolues: dict[str, str] = {}
    for colonne in colonnes:
        for alias in colonne.alias_normalises:
            if alias in presentes:
                resolues[colonne.cle] = presentes[alias]
                break

    manquantes = [c.libelle for c in colonnes if c.obligatoire and c.cle not in resolues]
    if manquantes:
        raise ImportColumnsMissingError(
            "Colonnes obligatoires absentes du fichier : " + ", ".join(manquantes) + ".",
            details={
                "colonnes_manquantes": manquantes,
                "colonnes_lues": [e for e in entetes if e],
                "colonnes_attendues": [c.libelle for c in colonnes],
            },
        )
    return resolues


# Signature d'une archive ZIP, dont un `.xlsx` est un cas particulier. Les deux
# variantes couvrent l'archive normale (`PK\x03\x04`) et l'archive vide
# (`PK\x05\x06`), qu'un tableur ne produit pas mais qu'un ZIP légitime peut être.
SIGNATURES_ZIP = (b"PK\x03\x04", b"PK\x05\x06")

# Extensions qui trahissent une intention claire, et permettent un message utile
# quand le contenu ne suit pas.
EXTENSIONS_PARLANTES = {
    ".csv": "un fichier texte CSV",
    ".xls": "un classeur Excel 97-2003",
    ".ods": "un classeur OpenDocument",
    ".txt": "un fichier texte",
    ".pdf": "un PDF",
}


def verifier_classeur(filename: str | None, contenu: bytes) -> None:
    """Écarte ce qui n'a aucune chance d'être un classeur, avant de le décompresser.

    Le verdict porte sur les **octets**, jamais sur le nom. C'est la règle que le
    projet applique déjà au logo d'organisation (ADR-020), et elle vaut ici dans
    les deux sens :

    - un fichier renommé `.xlsx` qui n'est pas une archive est refusé, quelle que
      soit son extension rassurante ;
    - un classeur parfaitement valide dont le **nom** a été abîmé en chemin est
      accepté. Le cas s'est présenté : un navigateur ou un client HTTP qui ignore
      `Content-Disposition` enregistre le modèle sous un nom arbitraire — vu en
      `export.csv` — et refuser sur cette base ferait rejeter notre propre modèle,
      intact, pour une raison qui ne regarde pas son contenu.
    """
    if not contenu:
        raise UnsupportedImportFileError("Le fichier envoyé est vide.")
    if len(contenu) > settings.max_upload_size_bytes:
        raise FileTooLargeError(
            details={
                "max_mb": settings.MAX_UPLOAD_SIZE_MB,
                "received_bytes": len(contenu),
            }
        )
    if contenu.startswith(SIGNATURES_ZIP):
        return

    nom = (filename or "").lower()
    reconnue = next(
        (libelle for ext, libelle in EXTENSIONS_PARLANTES.items() if nom.endswith(ext)),
        None,
    )
    detail = f"Le fichier envoyé est {reconnue}. " if reconnue else ""
    raise UnsupportedImportFileError(
        detail + "Seul le format `.xlsx` est accepté : réenregistrez votre fichier "
        "au format « Classeur Excel (.xlsx) », ou repartez du modèle téléchargeable, "
        "déjà au bon format. Un fichier `.xlsx` correct est accepté même si son nom "
        "a été modifié.",
        details={"filename": filename},
    )


def _ouvrir(contenu: bytes, *, data_only: bool) -> Workbook:
    try:
        return load_workbook(io.BytesIO(contenu), read_only=True, data_only=data_only)
    except (BadZipFile, InvalidFileException, KeyError, ValueError) as exc:
        raise UnsupportedImportFileError(
            "Le fichier n'est pas un classeur Excel `.xlsx` lisible. Le format "
            "`.xls` (Excel 97-2003) n'est pas accepté : réenregistrez-le au format "
            "`.xlsx`, ou repartez du modèle téléchargeable.",
            details={"raison": type(exc).__name__},
        ) from exc


def lire_lignes(
    contenu: bytes, colonnes: tuple[Colonne, ...], *, max_lignes: int
) -> tuple[list[dict[str, str]], int]:
    """Lit un `.xlsx` et rend `(lignes, nombre de lignes non vides)`.

    Chaque ligne est un `dict` indexé par `Colonne.cle`, valeurs déjà mises en
    texte et rognées, plus `__ligne__` — le numéro dans le classeur, seul repère
    qui permette à l'administrateur de retrouver sa cellule.

    Le classeur est ouvert **deux fois**, et ce n'est pas un gaspillage. Avec
    `data_only=True`, une cellule contenant une formule rend la valeur que le
    tableur a mise en cache — c'est ce qu'on veut, quelqu'un qui compose ses noms
    par `CONCATENER` doit pouvoir importer. Mais si le fichier n'a jamais été
    ouvert par un tableur, ce cache est vide et la cellule se lit **`None`** : la
    ligne partirait alors comme si elle était blanche, sans que rien ne le signale.
    La seconde lecture, formules brutes, sert uniquement à distinguer « vide » de
    « formule jamais calculée », et à le dire.

    `max_lignes` est un garde-fou, pas une pagination : au-delà, la lecture
    s'arrête et lève. Un import s'applique en une seule transaction, il n'a pas
    vocation à avaler cent mille lignes.
    """
    valeurs = _ouvrir(contenu, data_only=True)
    brut = _ouvrir(contenu, data_only=False)
    try:
        if not valeurs.worksheets:
            raise UnsupportedImportFileError("Le classeur ne contient aucune feuille.")
        rangees = valeurs.worksheets[0].iter_rows(values_only=True)
        rangees_brutes = brut.worksheets[0].iter_rows(values_only=True)

        entetes_bruts = next(rangees, None)
        next(rangees_brutes, None)
        if entetes_bruts is None:
            raise UnsupportedImportFileError("La première feuille est vide.")

        entetes = [cellule_en_texte(v) for v in entetes_bruts]
        resolues = resoudre_colonnes([e or None for e in entetes], colonnes)
        index = {cle: entetes.index(libelle) for cle, libelle in resolues.items()}

        lignes: list[dict[str, str]] = []
        formules: list[str] = []
        lues = 0
        for numero, (rangee, rangee_brute) in enumerate(
            zip(rangees, rangees_brutes, strict=False), start=2
        ):
            valeurs_ligne: dict[str, str] = {}
            for cle, position in index.items():
                cellule = rangee[position] if position < len(rangee) else None
                texte = cellule_en_texte(cellule)
                if not texte:
                    source = (
                        rangee_brute[position] if position < len(rangee_brute) else None
                    )
                    if isinstance(source, str) and source.startswith("="):
                        formules.append(f"{resolues[cle]} ligne {numero}")
                valeurs_ligne[cle] = texte

            if not any(valeurs_ligne.values()):
                continue
            lues += 1
            if lues > max_lignes:
                raise ImportTooManyRowsError(max_lignes)
            lignes.append({"__ligne__": str(numero), **valeurs_ligne})

        if formules:
            raise UncomputedFormulaError(formules)
        return lignes, lues
    finally:
        valeurs.close()
        brut.close()


# --- Écriture du modèle ------------------------------------------------------

_ENTETE_FOND = PatternFill("solid", fgColor="1F3864")
_ENTETE_POLICE = Font(color="FFFFFF", bold=True)
_TITRE = Font(bold=True, size=12)


def construire_modele(
    titre: str, colonnes: tuple[Colonne, ...], instructions: list[str]
) -> bytes:
    """Fabrique le classeur modèle : une feuille de saisie, une feuille de notice.

    La feuille de saisie ne contient **que** les en-têtes. Les exemples vivent sur
    la notice : laissés dans la grille, ils finiraient importés par le premier
    administrateur qui oublie de les effacer.

    Les colonnes sont forcées au format texte. Sans cela, Excel transforme le code
    « 007 » en nombre 7 dès la saisie, et l'import ne retrouve plus le service.
    """
    classeur = Workbook()

    saisie = classeur.active
    saisie.title = titre
    saisie.append([c.libelle for c in colonnes])
    for position, colonne in enumerate(colonnes, start=1):
        lettre = get_column_letter(position)
        cellule = saisie.cell(row=1, column=position)
        cellule.fill = _ENTETE_FOND
        cellule.font = _ENTETE_POLICE
        cellule.alignment = Alignment(horizontal="center", vertical="center")
        saisie.column_dimensions[lettre].width = colonne.largeur
        # Format texte porté par la **colonne**, et non par un millier de cellules
        # pré-créées : le style s'applique à tout ce que l'utilisateur saisira, sans
        # laisser mille lignes vides dans le fichier qu'il ouvre.
        saisie.column_dimensions[lettre].number_format = "@"
    saisie.freeze_panes = "A2"

    notice = classeur.create_sheet("Notice")
    notice.column_dimensions["A"].width = 26
    notice.column_dimensions["B"].width = 14
    notice.column_dimensions["C"].width = 72
    notice.column_dimensions["D"].width = 32

    notice.append(["Comment remplir ce modèle"])
    notice["A1"].font = _TITRE
    notice.append([])
    for instruction in instructions:
        notice.append([neutraliser_formule(instruction)])
    notice.append([])

    ligne_entete = notice.max_row + 1
    notice.append(["Colonne", "Obligatoire", "Description", "Exemple"])
    for position in range(1, 5):
        cellule = notice.cell(row=ligne_entete, column=position)
        cellule.fill = _ENTETE_FOND
        cellule.font = _ENTETE_POLICE
    for colonne in colonnes:
        notice.append(
            [
                colonne.libelle,
                "oui" if colonne.obligatoire else "non",
                neutraliser_formule(colonne.description),
                neutraliser_formule(colonne.exemple),
            ]
        )
    for rangee in notice.iter_rows(min_row=ligne_entete + 1):
        for cellule in rangee:
            cellule.alignment = Alignment(vertical="top", wrap_text=True)

    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()
